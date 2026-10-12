"""
Build a single Excel workbook, one row per GATE question, summarizing that
question plus every model's final answer and marks obtained -- built from the
merged data/responses/ and results/ data, which already carries the
official GATE marking including MCQ negative marking.

Buffer questions with no official answer key (EE_2025 #66-75) are excluded --
they were never part of the graded benchmark.

Columns:
    Paper Code, Year, Question Number, Section, Question Type,
    Question Format (Image/Text), Correct Answer,
    <Model> Final Answer  (x7),
    Total Marks,
    <Model> Marks Obtained  (x7)
"""

from __future__ import annotations

import glob
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from gatebench.config import ROOT
from gatebench.dataset import load_rows

OUT_PATH = ROOT / "reports" / "gate_benchmark_report.xlsx"

# folder name -> friendly display name, in the fixed column order to use.
MODEL_DISPLAY_NAMES = {
    "claude-4.7-opus-20260416": "Claude Opus 4.7",
    "gemini-3.1-pro-preview-20260219": "Gemini 3.1 Pro",
    "gpt-5.5-20260423": "GPT-5.5",
    "grok-4.3-20260430": "Grok 4.3",
    "kimi-k2.6-20260420": "Kimi K2.6",
    "mistral-large-2512": "Mistral Large",
    "qwen3.6-plus-04-02": "Qwen 3.6 Plus",
}


def format_correct_answer(answer) -> str:
    if answer is None:
        return ""
    if isinstance(answer, str):
        return answer  # voided (MTA) questions
    if isinstance(answer, list) and answer and isinstance(answer[0], list):
        # NAT: list of [low, high] pairs
        return " OR ".join(f"{lo} to {hi}" for lo, hi in answer)
    # MCQ / MSQ: list of option letters
    return ", ".join(answer)


def load_model_paper(model_dir: str, paper: str) -> dict[int, dict]:
    year = paper.rsplit("_", 1)[-1]
    rows = load_rows(model_dir, year, paper)
    return {row["question_number"]: row for row in rows}


def build():
    model_dirs = sorted(MODEL_DISPLAY_NAMES.keys())

    # Discover the paper list from any one model's directory tree.
    sample_model = model_dirs[0]
    from gatebench.dataset import iter_papers
    papers = [stem for _year, stem in iter_papers(sample_model)]

    per_model_data = {
        model_dir: {paper: load_model_paper(model_dir, paper) for paper in papers}
        for model_dir in model_dirs
    }

    wb = Workbook()
    ws = wb.active
    ws.title = "GATE Benchmark Results"

    headers = [
        "Paper Code",
        "Year",
        "Question Number",
        "Section",
        "Question Type",
        "Question Format",
        "Correct Answer",
    ]
    headers += [f"{MODEL_DISPLAY_NAMES[m]} Final Answer" for m in model_dirs]
    headers.append("Total Marks")
    headers += [f"{MODEL_DISPLAY_NAMES[m]} Marks Obtained" for m in model_dirs]

    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for paper in papers:
        code, year = paper.rsplit("_", 1)
        # Use the first model's rows as the canonical question list/order --
        # all models' JSONs mirror data/questions/ for the same paper.
        canonical_rows = load_rows(sample_model, year, paper)
        for row in canonical_rows:
            if row["question_type"] is None:
                continue  # buffer question with no official key -- skip

            qnum = row["question_number"]
            image_based = "Image" if row.get("has_image") else "Text"
            correct_answer = format_correct_answer(row["correct_answer"])

            excel_row = [
                code,
                int(year),
                qnum,
                row["section"],
                row["question_type"],
                image_based,
                correct_answer,
            ]
            for model_dir in model_dirs:
                m_row = per_model_data[model_dir][paper].get(qnum)
                excel_row.append(m_row["final_answer"] if m_row else None)

            excel_row.append(row["marks"])
            for model_dir in model_dirs:
                m_row = per_model_data[model_dir][paper].get(qnum)
                excel_row.append(m_row["marks_obtained"] if m_row else None)

            ws.append(excel_row)

    ws.freeze_panes = "A2"
    for i, _ in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = 16

    OUT_PATH.parent.mkdir(exist_ok=True)
    wb.save(OUT_PATH)
    print(f"Wrote {ws.max_row - 1} question rows to {OUT_PATH}")


if __name__ == "__main__":
    build()
